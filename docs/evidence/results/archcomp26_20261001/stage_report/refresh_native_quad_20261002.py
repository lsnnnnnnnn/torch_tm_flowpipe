"""Refresh the stage DOCX from saved QUAD receipts; launch no experiments."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


STAGE = Path(__file__).resolve().parent
ROOT = STAGE.parents[4]
RESULTS = STAGE.parent
DOCX = STAGE / "ARCHCOMP26_STAGE_REPORT_DRAFT_20261001.docx"
NATIVE = RESULTS / "native_quad_paper_full50_001"
FIGURE = RESULTS / "quad_paper_fourway_saved_20261002/quad_paper_fourway_t_x3_pooled_tube.png"


def one(doc, prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, (prefix, len(matches))
    return matches[0]


def main():
    scan = json.loads((NATIVE / "SCAN.json").read_text())
    result = json.loads((NATIVE / "RESULT.json").read_text())
    assert scan["record_count"] == scan["expected_record_count"] == 1_024_000
    assert scan["complete_unique_ordered_grid"] and not any(scan["errors"].values())
    assert scan["terminal_endpoint_x3_within_0_94_1_06"]
    assert result["status"] == "completed" and result["exit_code"] == 0 and not result["timed_out"]
    assert FIGURE.is_file()
    subprocess.run([sys.executable, str(ROOT / "tools/revise_archcomp26_stage_report_20261001.py")], check=True)
    doc = Document(DOCX)
    x3 = scan["terminal"][2]
    lo, hi = x3["endpoint_lo"], x3["endpoint_hi"]
    wall = result["wall_s"]

    p = one(doc, "目标是在同一 benchmark 合同下")
    p.text = p.text.replace("103 条新尝试，25 个", "105 条新尝试，26 个")
    p.text = p.text.replace("其中 DP less、ACC", "其中 QUAD、DP less、ACC")
    p = one(doc, "Double Pendulum less-robust 的新 Flow*")
    p.text = (
        "Double Pendulum less-robust 的新 Flow* native、Huan、Xiangru 均覆盖 225 初盒、20 控制期和 100 个 ODE 子步；单次进程 wall 分别为 1107.127423、9.539174、8.387790 s。保存 tube 在连续全时域安全带内。"
        f"论文方程 QUAD 四方均有 1024 初盒、50 期、1000 子步的单次完整数值运行；P3/Huan/Xiangru/native 的进程 wall 分别为 1357.555、94.583、108.018、{wall:.3f} s。"
        "四方 T=5 的 x₃ endpoint 均在目标带内，作者 checker 均打印 VERIFIED；这不是独立端到端浮点证明，也不构成稳定速度排名。"
    )
    p = one(doc, "两物理态 Single Pendulum 的四方")
    p.text = p.text.replace(
        "DP more 三方仍未完成 T=0.4。",
        "DP more 四方均未完成 T=0.4，P3 仅有全 225 盒首周期数值前缀且性质 Unknown。",
    ).replace(
        "修正官方 unsafe 盒后的 Attitude",
        "TORA reach-sigmoid 的 Huan 官方 u=11f 配置仅有完整初盒一期前缀、终点性质未检；修正官方 unsafe 盒后的 Attitude",
    )
    one(doc, "6  论文方程 QUAD").text = "6  论文方程 QUAD · 新四方完整数值时域"
    one(doc, "P3、Huan、Xiangru 各完成").text = (
        "P3、Huan、Xiangru 与 Flow* native 各完成 1024 盒 × 1000 小步的单次全程。"
        f"四方外层 wall 依次为 1357.555 / 94.583 / 108.018 / {wall:.3f} s。"
        "P3 1,024,000/1,024,000 盒步接受；native ranges.bin 独立重扫有 1,024,000 条唯一有序记录，数值有限、区间非逆序、endpoint 均含于同一步 tube。"
        f"native T=5 x₃ endpoint union 为 [{lo},{hi}]；其原始外层 RESULT 为 completed/exit 0，作者日志打印 VERIFIED。"
    )
    one(doc, "P3 为最新严格数值引擎诊断").text = (
        "P3 的 T=5 x₃ endpoint union 为 [0.9584732146312492,1.0256989633476477]，宽 0.06722574871639841；"
        "Huan/Xiangru P2 parity 同为 [0.967434441417146,1.015176258384569]，宽 0.04774181696742297；"
        f"native 为 [{lo},{hi}]，宽 {hi-lo:.16g}。"
        "P3 的 end_to_end_strict_certificate=false；保存盒、作者 verdict 和较窄宽度都不是独立端到端证明。"
        "native 保存源码只在终时逐盒检查目标，尚无论文 reach-and-remain 的全时证明。"
        "四方数值引擎、资源与计时路径不同，各只有一条完整进程，不给速度名次。旧约 75–80 s Huan 属另一动力学合同。"
    )
    one(doc, "6.1  P3 QUAD").text = "6.1  QUAD native 与 P3 全时 tube 及两方终点"
    one(doc, "每个局部小步只有 1024 个").text = (
        "native 与 P3 的图线来自各自 1000 个小步的 1024 盒 x₃ pooled tube；Huan/Xiangru 在本次保存证据中仅有 T=5 endpoint 聚合，图中只标终点，不构造其全时曲线。"
        "[0.94,1.06] 仅是 T=5 终点目标；native 的 box 投影不恢复 octagon 坐标相关性。"
    )
    caption = one(doc, "图 3  新论文方程 QUAD P3")
    caption_index = next(i for i, p in enumerate(doc.paragraphs) if p.text == caption.text)
    picture = doc.paragraphs[caption_index - 1]
    assert picture._p.xpath(".//w:drawing")
    picture.text = ""
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(FIGURE), width=Cm(16.8))
    caption.text = "图 3  论文方程 QUAD：native/P3 完整保存 t–x₃ tube，Huan/Xiangru 只标 T=5 endpoint。"
    one(doc, "最后一行 observer 的 x₃ endpoint").text = (
        "四方都完成目标数值时域，但只有 native/P3 的逐步 tube 能从现存保存区间重画。"
        "图中 P3 终点来自 observer 保存行 [0.9584732146312499,1.0256989643984427]，表中采用驱动最终 hull；两者不可混写。"
        "图源及原生范围扫描见 quad_paper_fourway_saved_20261002/SUMMARY.md 与 native_quad_paper_full50_001/SCAN.json；MATLAB 脚本未实跑。"
    )
    one(doc, "论文方程 QUAD：Flow* native 全程尚在运行").text = (
        "论文方程 QUAD 四方已有单次完整数值时域，native 原始 RESULT 为 completed/exit 0、作者 checker 为 VERIFIED；"
        "重复计时和独立端到端证明仍缺。现有四方图只提供 native/P3 全时 tube 及 Huan/Xiangru 终点。"
    )
    p = one(doc, "Huan/Xiangru 两条早停运行")
    p.text += (
        " P3 另有全 225 盒首个 0.02 s 控制期的 4/4 小步、900/900 盒步数值接受；"
        "第 3/4 步部分保存区间越过安全带，作者驱动报 Unknown，余下 19 期未尝试。"
    )
    one(doc, "DP more：核查早停区间").text = (
        "DP more：四方都缺 T=0.4 全程；保留 native 65–80 步、Huan/Xiangru 73–80 步和 P3 第 2–20 期为未观察，"
        "并分别核查早停区间与 P3 首周期 Unknown 的性质语义。"
    )
    one(doc, "“全程”表示本轮一个方法").text = (
        "“全程”只表示目标数值时域全部推进，不自动证明性质；“诊断”是短前缀，“早停”或“未知”没有完整时域。"
        "QUAD、DP less、ACC、两物理态 Single Pendulum、修正后 Attitude、Docking 各有四方全程，TORA remain 原生/P3 也全程，共 26 格。"
        "Docking 原生外层 failed/exit 2 源于 checker Unknown，数值仍完成 400 段；ACC 的 24 次独立进程 campaign 不构成稳定速度排名。"
        "TORA Huan/Xiangru 虽观察到第 200 小步，但部分盒步拒绝，故列为未知前缀。"
    )

    for row in doc.tables[0].rows:
        if row.cells[0].text.startswith("QUAD 论文方程"):
            row.cells[0].text = "QUAD 论文方程 · 四方 · 新 2026"
            row.cells[1].text = "各 1024 盒 × 1000 子步；50 控制期"
            row.cells[2].text = "四方单次全程；native 1,024,000 条范围重扫齐全，作者 VERIFIED；P3 严格证书 false。无稳定排名。"
        if row.cells[0].text.startswith("DP more · 三种方法"):
            row.cells[0].text = "DP more · 四方 · 新 2026"
            row.cells[1].text = "native 64/80；Huan/Xiangru 各 72/80；P3 全 225 盒首周期"
            row.cells[2].text = "三方早停；P3 900/900 盒步接受但性质 Unknown。四方均无 T=0.4 全程结果。"
    quad = doc.tables[3]
    quad.rows[0].cells[1].text = "P3 / Huan / Xiangru / native 新全程记录"
    quad.rows[2].cells[1].text = f"outer wall 1357.555 / 94.583 / 108.018 / {wall:.3f} s；各单次进程"
    quad.rows[3].cells[1].text = "四方 1024×1000 全程；P3 全部盒步接受；native 1,024,000 条完整范围"
    quad.rows[4].cells[1].text = (
        "P3 [0.9584732146312492,1.0256989633476477]；H/X [0.967434441417146,1.015176258384569]；"
        f"native [{lo},{hi}]"
    )
    quad.rows[5].cells[1].text = "目标 [0.94,1.06]；四方作者 checker VERIFIED；无独立端到端证书"
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text == "quad-reach" and len(row.cells) == 5:
                row.cells[4].text = "全程"
            if row.cells[0].text == "double-pendulum-more-robust" and len(row.cells) == 5:
                row.cells[1].text = "首周期·性质未知"
                row.cells[4].text = "早停·性质未知"
            if row.cells[0].text == "tora-reach-sigmoid" and len(row.cells) == 5:
                row.cells[2].text = "全初盒一期"

    assert len(doc.inline_shapes) == 6
    assert "105 条新尝试，26 个" in one(doc, "目标是在同一 benchmark 合同下").text
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
