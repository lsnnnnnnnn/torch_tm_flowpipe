#!/usr/bin/env python3
"""Assemble current reports from saved evidence only; no solver, checker or digest."""
from contextlib import ExitStack
import argparse, csv, hashlib, importlib.util, json, os, re
from pathlib import Path
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/evidence/results/archcomp26_report_20261006'
OLD=ROOT/'docs/evidence/results/archcomp26_report_20261005'
STUDY=ROOT/'research/p3_speed_tightness_20261006'
SOURCE=ROOT/'docs/ARCHCOMP26_FINAL_REPORT_DRAFT.md'
METHODS=('pytorch_gpu','huan','xiangru','flowstar_native')
LABELS={'pytorch_gpu':'P3','huan':'Huan','xiangru':'Xiangru','flowstar_native':'原生 Flow*'}

def read(path): return json.loads(path.read_text())
def f(x): return '—' if x is None else f'{x:.8g}'
def sec(x): return '—' if x is None else f'{x:.3f}'
def rel(path): return os.path.relpath(ROOT/path, SOURCE.parent)
def link(label,path): return f'[{label}]({rel(path)})'
def table(headers,rows):
    def clean(v): return str(v).replace('|','/').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |',*['| '+' | '.join(map(clean,row))+' |' for row in rows]])+'\n'

def source_line(cells):
    return '原始结果：'+'；'.join(link(LABELS[c['method']], c['current_selected_receipts'][0]['path']) if c['current_selected_receipts'] else LABELS[c['method']]+'未启动' for c in cells)+'。'

def load_builder():
    path=ROOT/'docs/evidence/results/archcomp26_20261001/stage_report/build_archcomp26_report_20261004.py'
    spec=importlib.util.spec_from_file_location('archived_docx_helpers',path)
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    module.SOURCE=SOURCE;module.OUT=OUT/'ARCHCOMP26_REPORT_20261006.docx'
    return module

def docx_build():
    h=load_builder()
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor
    from docx.package import ImageParts
    def addtable(doc,block):
        rows=[h.table_cells(x) for x in block];assert h.is_divider(rows[1])
        rows=[rows[0],*rows[2:]];cols=len(rows[0]);assert 2<=cols<=5
        widths={2:[4.0,13.5],3:[4.5,6.5,6.5],4:[5.2,4.1,4.1,4.1],5:[4.5,3.25,3.25,3.25,3.25]}[cols]
        t=doc.add_table(rows=0,cols=cols);t.autofit=False;t.style='Table Grid'
        for ridx,row in enumerate(rows):
            assert len(row)==cols,(len(row),cols,row)
            cells=t.add_row().cells
            for j,value in enumerate(row):
                c=cells[j];c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                h.shade(c,'E8EDF1' if ridx==0 else ('F7F8FA' if ridx%2==0 else 'FFFFFF'))
                p=c.paragraphs[0];p.paragraph_format.keep_with_next=(ridx==0 or (len(rows)<=9 and ridx<len(rows)-1));p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.08
                if j: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                h.inline(p,value)
                for r in p.runs:r.font.size=Pt(9.3);r.bold=ridx==0
            h.avoid_row_split(t.rows[ridx])
        h.repeat_header(t.rows[0]);p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=0.3
    with ExitStack() as guard:
        for name in (*hashlib.algorithms_guaranteed,'new','file_digest'):guard.enter_context(patch.object(hashlib,name,h.forbid_digest))
        guard.enter_context(patch.object(ImageParts,'get_or_add_image_part',h.image_part_by_bytes))
        doc=Document();h.initialize(doc)
        for name,size in [('Normal',10.5),('Stage Note',10),('Evidence Body',10),('List Bullet',10.5),('Caption',9),('Title',20),('Heading 1',14),('Heading 2',11.5)]:
            doc.styles[name].font.size=Pt(size)
        doc.styles['Normal'].paragraph_format.line_spacing=1.17
        doc.styles['Normal'].paragraph_format.space_after=Pt(5)
        doc.sections[0].footer.paragraphs[0].clear()
        foot=doc.sections[0].footer.paragraphs[0];foot.text='2026 年 10 月 6 日  ·  时间和保存宽度比较  ·  ';foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;h.add_page_field(foot)
        for r in foot.runs:h.set_font(r.font,r._element);r.font.size=Pt(8)
        doc.core_properties.title='ARCH COMP26 四方时间与区间宽度比较报告'
        lines=SOURCE.read_text().splitlines();i=0
        while i<len(lines):
            line=lines[i]
            if not line.strip():i+=1;continue
            if line.startswith('|'):
                block=[]
                while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
                addtable(doc,block);continue
            if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
            elif line.startswith('## '):
                p=doc.add_heading(line[3:],level=1)
                if line.startswith('## 附录 A'):p.paragraph_format.page_break_before=True
            elif line.startswith('### '):doc.add_heading(line[4:],level=2)
            elif line.startswith('![矢'):
                pass
            elif line.startswith('!['):
                match=h.IMAGE_LINK.search(line)
                if match:h.add_figure(doc,*match.groups())
            elif line.startswith('- '):h.add_text(doc,line[2:],'List Bullet')
            elif line.startswith('> '):h.add_text(doc,line[2:],'Stage Note')
            else:
                p=h.add_text(doc,line)
                if line.startswith(('**共同终点','**截至共同','**各法自身')):p.paragraph_format.keep_with_next=True
            i+=1
        doc.save(h.OUT)
    return h

def assemble():
    metadata=read(OLD/'report_metadata.json'); status=read(OUT/'blockers/status_64cells.json'); timing=read(OUT/'timing/timing_index.json'); width=read(OUT/'widths/summary.json'); runs=read(STUDY/'report_data/current/RUN_INDEX.json')['runs']
    selection=read(STUDY/'report_data/selection.json'); audit=read(STUDY/'report_data/current/selection_audit.json')
    ws={(r['instance_id'],r['method'],r['state']):r for r in width}; cells={(r['instance_id'],r['method']):r for r in status['cells']}
    specs={r['instance_id']:r for r in read(OLD/'widths/instances.json')}
    def chosen(key,method):
        runkey=timing['benchmarks'][key]['methods'][method].get('current_selected_run')
        return timing['runs'].get(runkey) if runkey else None
    def timer(key,method,layer):
        r=chosen(key,method);return r['timings'].get(layer) if r else None
    lines=[]
    def add(value=''):lines.append(value)
    add('# ARCH COMP26 速度优化与逐状态宽度比较报告')
    add('2026 年 10 月 6 日北京时间更新　16 个 benchmark × 4 种方法　保存的完整时间和范围证据')
    add('本轮先减少重复计算，再测试区间收紧，已有多个完整实例的单次求解时间下降。QUAD 的联合方案在同一次完整 1024 盒 × 1000 步运行中，把 process 从 1004.197 秒降到 978.067 秒（少 2.60%），高度 x3 终点宽度缩小 2.54%，x11 缩小 4.76%；x9 仍有 8 条极小变宽，不能称全程处处改善。TORA 进一步收紧的收益较小，且比只提速的方案略慢。QUAD 仍未达到 Huan / Xiangru 的完整进程速度或高度终点紧度。')
    add('当前报告按实际采用的结果更新全部时间、逐状态 endpoint / tube 宽度、每个未完成方法的停点和原因。当前矩阵仍为 38 个项目内数值完整、8 个同合同历史完整、14 个无完整数值时域、4 个 Airplane discrete 合同阻塞，共 64 格。新候选改进已有格，不增加覆盖数；旧实验、旧资格检查器和原始 RESULT 未重跑或改写。图形仍全部由 Python 生成。')
    add('时间单位为秒，宽度为保存上界减下界，按物理状态原单位列出。宽度显示八位有效数字，完整可读 binary64、上下界、逐步范围和原始路径在 CSV / JSON。不同量纲不合成紧度总分。“—”是缺失或不适用，不表示零。数值完成、保存的性质观察、集合包含和独立端到端 NNCS 浮点证书分开报告。')
    add('表中 P3 是我方实现族的沿用列名，不代表所有实例阶数和注入路径相同。具名两态 Single Pendulum 保留 order2 / point1 / validation3、native-f64 和原参考注入，仅使用其已安装的严格 endpoint；其他 P3 合同按各自保存配置说明。')
    add('## 本轮完整候选的速度与紧度取舍')
    add('所选方案与上一版采用方案的内部 driver 计时见下表；每个新方案仅一次实际运行。旧分支的最快中位数使用不同合同、保留阶数或输出工作量，不能直接移入此表。共享服务器和不同时间的单次测量不能证明稳定加速比。')
    rows=[]
    for row in audit:
        key=row['instance_id'] if 'instance_id' in row else row['benchmark']
        meta=next(x for x in metadata if x['instance_id']==key)
        before=row['old_timings']['driver_elapsed_s']; after=row['new_timings']['driver_elapsed_s']
        policy=selection[key]['width_policy']
        rows.append([meta['title'],sec(before)+' → '+sec(after),f'{100*(before-after)/before:.2f}%' if before and after else '—',sec(row['new_timings']['process_wall_s']),'新保存范围' if policy=='new_saved_ranges' else '比较输出相同'])
    add(table(['实例','driver 上一版 → 本轮','单次时间下降','本轮 process','宽度依据'],rows))
    add('![本轮所选方案与上一版内部时间](evidence/results/archcomp26_report_20261006/figures/driver_selected_before_after.png)')
    add('TORA tanh 的速度优先版 process / driver 为 8.791 / 4.563 秒；收紧版为 9.242 / 4.963 秒。收紧版四态 500 步 endpoint 和 tube 的 4,000 项宽度均不增加，终点四态仅缩小约 0.023%–0.048%；部分终点区间位置变化，因此更窄不等于都被旧区间包含。')
    add('TORA sigmoid 的上一版速度优先结果为 9.392 / 5.255 秒。仅把 cutoff 从 1e-6 改到 1e-8 的完整结果为 9.995 / 5.763 秒，x1 / x2 终点缩小约 0.86% / 0.91%，全 4,000 项保存宽度均不增加且各区间包含于旧区间。高阶候选及其早期变宽项另列于附录，不能凭最后两态的改善声称全程更紧。')
    add('## 全部 benchmark 的完整进程时间')
    add('仅对覆盖具名全时域的方法填写明确保存的 process wall；历史复用注明“历史”。失败或短前缀耗时列入分节和完整时间索引，不参加全程速度比较。外层进程、wrapper、payload、driver 内部和 driver 调用分别计时，不互相代填。')
    time_rows=[]
    for meta in metadata:
        key=meta['instance_id']; row=[meta['title']]
        for m in METHODS:
            c=cells[key,m]; v=timer(key,m,'process_wall_s')
            row.append(sec(v)+(' 历史' if c['coverage_classification']=='historical_same_contract_numeric_full' and v is not None else '') if c['numeric_progress']['full_numeric_horizon'] else '未全程' if c['status']!='contract_blocked' else '合同缺失')
        time_rows.append(row)
    add(table(['实例','P3','Huan','Xiangru','原生'],time_rows))
    add('NAV 与 TORA tanh 部分历史作者记录没有外层 process，只有 driver call 或内部 elapsed，故主表保留空缺。NAV standard 新方案 process 24.031 秒不能直接对比作者约 13 秒的内部计时；相同层级的完整数据在分节展开。')
    add('### 既有重复进程计时')
    add('下面是冻结 campaign 后五个独立进程的中位数 [最小值, 最大值]，每格 n=5，不包含首进程，也不混入本轮单次优化。新旧 campaign 不互相替代。')
    rows=[]
    for meta in metadata:
        key=meta['instance_id']; methods=timing['benchmarks'][key]['methods']
        if not methods['pytorch_gpu'].get('campaign'): continue
        row=[meta['title']]
        for m in METHODS:
            d=methods[m]['campaign']['later_statistics']['process_wall_s']; row.append(f"{sec(d['median'])} [{sec(d['min'])}, {sec(d['max'])}]")
        rows.append(row)
    add(table(['冻结 campaign','P3','Huan','Xiangru','原生'],rows))
    add('### 学到和采用的实现变化')
    add('旧 codex/progress-report-20260923 分支中的 engine_linear_leaf_v2 路径已经在现有引擎中，并没有丢失。它当时较快的结果包含不同控制定义、非同等级注入或不保存逐步几何的配置；不能取消现有严格路径和绘图输出来复现那个数字。上一版恢复了私有输出、Horner 绑定和按盒数分块，本轮把原来两轮加权映射放在同一 CUDA 图内，并保留完整条件检查、失败回退和首个真实输入的原实现直接比对。')
    add('分块不是越小越快。NAV robust 25 盒用 32 行融合图有收益；NAV standard 640 盒用 256 行分块变慢，512 行融合图才把内部时间降到 19.864 秒。QUAD 1024 盒的融合版没有胜过已选 256 行路径；相同输入 sin/cos 幂复用的全程结果为 process 975.862 / driver 968.516 秒，相比上一版 1004.197 / 997.675 秒分别少 2.82% / 2.92%。1000 条 pooled 观察直接相同。新 GPU 小门的 0.938 秒计入外层，不在 driver 内。最终主选把这个复用与控制余项收紧组合，process / driver 为 978.067 / 969.549 秒，分别比上一版少 2.60% / 2.82%；相比纯提速版付出约 2.205 秒 process，换取自身记录的终点收紧。新联合 GPU 小门为 0.934 秒，仍计在外层。')
    add('Huan / Xiangru 的 QUAD 使用 work / point / validation 阶 2/1/1 与 parity 路径，P3 原选方案为 3/2/4 且保存严格区间误差账本；双方工作量并不相同。本轮借鉴其减少重复工作和更高阶 TORA 配置，但没有移除严格验证。K=20 是每 20 步重算完整保留历史，不是截掉 20 步以前的历史。未做匹配消融，不能给这些差异编造因果百分比。')
    add('控制余项收紧借鉴作者的 hybrid / 双斜率仿射包络思路，同时保留我方原控制多项式与严格注入：把额外包络换算成相对于这个固定多项式的余项约束，再求交。QUAD 全程实际需要 50 次原 NN 调用加 50 次额外调用；原计数器的 50 不能冒充总数。附加浮点 CROWN 包络仍是条件性输入，不因此获得独立端到端证书。')
    add('## 宽度比较口径')
    add('每步先对相同方法的全部有效初盒取坐标并集，endpoint 是传播终点，tube 是整个小步。正文先列四方共同有效终点的宽度，再列截至该时刻每个单步 tube 宽度的最大值；后者不是全时间并集的宽度。辅助时钟、保持控制及常值扰动不加入物理态评分，原记录保留。只有实际保存过相关方向或逐盒数据才可分析该几何；pooled 坐标盒不能恢复相关八方向包络。')
    add('提前停止的共同数值时刻与安全观察时刻不同：Balancing raw4 为 0.415 秒；DP more 数值 0.32 秒、安全 0.30 秒；TORA remain h=0.1 数值 18.9 秒、安全 18.4 秒。各法自身最后完整初集范围另列，不能用不同时刻的宽度排名。QUAD 作者两法只存终点，全时 tube 缺失。')
    add('新数值参数的 TORA 宽度从它们自己的 ranges.bin 重算，QUAD 联合方案从自身逐步观察重算；仅在完整保存输出直接比较相同的实现候选上沿用原范围。时间来源与几何来源分别注明，不伪装为同一原始运行。逐态绝对差、相对差和终点包含关系见配套 pairwise_comparisons.csv。')
    comparisons=[]
    for meta in metadata:
        key=meta['instance_id'];cs=[cells[key,m] for m in METHODS];spec=specs[key]
        add(f"## {meta['number']} {meta['title']}")
        add(meta['contract'])
        add('### 数值范围与未完成原因')
        rows=[]
        for c in cs:
            p=c['numeric_progress'];full=p['full_numeric_horizon'];prefix=p['full_initial_set_accepted_or_valid_saved_prefix_steps']
            rows.append([LABELS[c['method']], '完整' if full else '合同缺失' if c['status']=='contract_blocked' else '无全程', f"{prefix if prefix is not None else '—'} / {p['planned_substeps_or_transitions']}", f(p['full_initial_set_prefix_time_s']),('性质窗见正文' if key=='single-pendulum-reach' else '不适用' if key in ('quad-reach','tora-reach-sigmoid','tora-reach-tanh','unicycle-reach') else f(p['saved_all_boxes_safe_prefix_time_s']))])
        add(table(['方法','数值状态','全初集步数','数值前缀 t','保存安全前缀 t'],rows))
        grouped={}
        for c in cs:
            text=c['reason_zh']+' '+c['needed_evidence_or_change_zh'];grouped.setdefault(text,[]).append(LABELS[c['method']])
        for text,labels in grouped.items(): add('**'+' / '.join(labels)+'：** '+text)
        verdicts={}
        for c in cs:verdicts.setdefault(c['verdict']['saved_geometry_observation_zh'],[]).append(LABELS[c['method']])
        for text,labels in verdicts.items():add('保存性质观察（'+' / '.join(labels)+'）：'+text)
        add(source_line(cs))
        if key in selection:
            selected=selection[key]
            add('当前 P3 使用 '+link(selected['run_id'],f"research/p3_speed_tightness_20261006/results/{selected['run_id']}/RESULT.json")+'；'+('宽度来自此候选的新保存范围。' if selected['width_policy']=='new_saved_ranges' else '宽度来自原保存参考，候选对完整保存对象有直接比较收据。'))
        add('### 所选进程各层时间')
        add('下表为各方法所选单次；未全程方法的数值只是这次失败或早停的耗时，不能参与完整运行速度比较。缺失表示该层没有独立保存，不能从另一层代填。')
        layers={'process_wall_s':'外层 process','wrapper_wall_s':'候选 wrapper','payload_wall_s':'runner payload','driver_elapsed_s':'内部 driver','driver_call_wall_s':'调用 driver','native_process_wall_s':'原生子进程','server_startup_s':'控制器启动','watchdog_process_wall_s':'watchdog'}
        rows=[]
        for field,label in layers.items():
            vals=[timer(key,m,field) for m in METHODS]
            if any(v is not None for v in vals):rows.append([label,*[sec(v) for v in vals]])
        if rows:add(table(['计时层级','P3','Huan','Xiangru','原生'],rows))
        elif key=='airplane-discrete':add('四方法均未在权威离散合同下启动，因此没有运行耗时或资源记录；另立的 CPU 诊断不填方法格。')
        else:add('没有具备已确认计时边界的所选进程。已保存入口尝试和未分类的原始 index wall 保留于逐进程 CSV，未猜测成求解时间。')
        res=[]
        for m in METHODS:
            r=chosen(key,m)
            if not r:continue
            rr=r['resources'];res.append(LABELS[m]+f" CPU={rr.get('cpu_affinity') if rr.get('cpu_affinity') is not None else '未记录'}，GPU={rr.get('physical_gpu') if rr.get('physical_gpu') is not None else '无或未记录'}")
        if res:add('保存资源：'+'；'.join(res)+'。进程起止、同批并发窗口、字段路径详见时间索引。')
        add('### 全部物理状态宽度')
        states=spec['physical_states'];ctimes=sorted({ws[key,m,s]['common_time'] for m in METHODS for s in states if ws[key,m,s]['common_time'] is not None})
        if ctimes:add('四方共同数值终点 t='+', '.join(f(x) for x in ctimes)+' 秒。若某一状态没有此共同终点，保持空白。')
        else:add('主合同不存在四方共同有效数值终点，以下空栏标明数据缺失，不借用其他合同或分盒短诊断。')
        if key.startswith('airplane-'):
            add('全部 12 个状态 '+', '.join(states)+'：四方法均无主合同有效 endpoint / tube 宽度，不能填终点或全时数值。逐状态缺失格仍全部列在 summary.csv 和 missing_fields.csv。')
        else:
            add('**共同终点 endpoint 绝对宽度**')
            add(table(['状态','P3','Huan','Xiangru','原生'],[[s,*[f(ws[key,m,s]['common_endpoint_width']) for m in METHODS]] for s in states]))
            add('**截至共同数值时刻的最大单步 tube 绝对宽度**')
            add(table(['状态','P3','Huan','Xiangru','原生'],[[s,*[f(ws[key,m,s]['common_max_tube_width']) for m in METHODS]] for s in states]))
        for other in METHODS[1:]:
            narrow=[];wide=[];equal=[]
            for s in states:
                a=ws[key,'pytorch_gpu',s];b=ws[key,other,s]
                for field in ('common_endpoint_width','common_max_tube_width'):
                    x,y=a[field],b[field]
                    if x is None or y is None:continue
                    contained=(a['common_endpoint_lo']>=b['common_endpoint_lo'] and a['common_endpoint_hi']<=b['common_endpoint_hi']) if field=='common_endpoint_width' else None
                    comparisons.append(dict(instance_id=key,state=s,geometry=field,reference_method=other,common_time=a['common_time'],p3_width=x,reference_width=y,difference=x-y,relative_difference_percent=100*(x/y-1) if y else None,p3_endpoint_subset=contained))
                x,y=a['common_endpoint_width'],b['common_endpoint_width']
                if x is None or y is None:continue
                (narrow if x<y else wide if x>y else equal).append(s)
            if narrow or wide or equal:add('共同终点 P3 对 '+LABELS[other]+'：'+('较窄 '+', '.join(narrow)+'；' if narrow else '')+('较宽 '+', '.join(wide)+'；' if wide else '')+('相等 '+', '.join(equal)+'；' if equal else '')+'逐态差值见配套 CSV。该比较不外推到其他时刻或未保存相关方向。')
        if not all(c['numeric_progress']['full_numeric_horizon'] for c in cs) and ctimes:
            add('**各法自身最后完整初集终点宽度**（时刻不同，仅记录，不横向排名）')
            add(table(['状态','P3','Huan','Xiangru','原生'],[[s,*[f(ws[key,m,s]['own_last_complete_endpoint_width'])+' @'+f(ws[key,m,s]['own_last_complete_t']) for m in METHODS]] for s in states]))
        if key=='quad-reach':
            add('P3 主表使用观察器最后 endpoint；driver final_hull 是另一个保存对象，其坐标宽度不可互换。Huan / Xiangru 只有 driver 终态坐标范围，逐步 tube 未保存，因此不能比较其全时紧度。原生 SCAN.json 保存了完整 1000 步 × 12 态 pooled 界，本版直接读取已保存扫描，没有重新扫描远端大文件。')
            add('本版纠正一处旧文字口径：旧报告约 5.69e-6 指 x5 的单侧界差，实际 driver 与 pooled 宽度差约 −1.1272e-5。新数值候选的差值另从自身两个保存对象计算，逐态列在 geometry_observations.json；原上下界记录不改。')
            add('原生历史终点 VERIFIED 不是参与者 reach-and-remain 全时 checker。高度带后缀只作为保存几何观察，需联合 tube 与 endpoint；它不补齐 Huan / Xiangru 的逐步数据或独立全时证书。当前 P3 的相关数值应从其自身新保存范围读取，不能继承旧轨迹的进入时刻。')
            if key in selection:
                detail_path=OUT/'widths'/(selection[key]['run_id']+'_geometry_observations.json')
                if detail_path.is_file():
                    detail=read(detail_path); suffix=detail['height_band']['suffixes']['tube_endpoint_union']
                    add('当前保存的 x3 tube 与 endpoint 联合范围，从第 '+str(suffix['first_saved_step'])+' 步（名义段起点 '+f(suffix['nominal_segment_start_s'])+' 秒）开始，后续全部处于 [0.94,1.06]；这是数据观察，不升级为全时性质证明。'+link('逐态保存对象差异及后缀来源',str(detail_path.relative_to(ROOT)))+'。')
            if selection.get(key,{}).get('width_policy')=='new_saved_ranges':
                oldw={(r['instance_id'],r['method'],r['state']):r for r in read(OLD/'widths/summary.json')}
                a=oldw[key,'pytorch_gpu','x3']['common_endpoint_width']; b=ws[key,'pytorch_gpu','x3']['common_endpoint_width']
                add(f'本次主选在同一次完整运行中记录当前时间与新宽度；高度 x3 终点从 {f(a)} 缩到 {f(b)}，比上一版窄 {100*(a-b)/a:.4f}%。各态全程细节、微小反向变化及速度优先备选见附录，不声称所有状态每一时刻都严格改善。')
                add('![QUAD 当前保存时间高度范围](evidence/results/archcomp26_report_20261006/figures/quad_time_x3.png)')
                add('![QUAD 当前轴对齐状态投影](evidence/results/archcomp26_report_20261006/figures/quad_x1_x2.png)')
                add('全部 12 态逐步宽度图提供可缩放 PDF：'+link('endpoint','docs/evidence/results/archcomp26_report_20261006/figures/quad_endpoint_widths.pdf')+'；'+link('tube','docs/evidence/results/archcomp26_report_20261006/figures/quad_tube_widths.pdf')+'。Huan / Xiangru 仅画真实保存的终态，缺失时段留空；坐标盒投影不代表相关八方向包络。')
            add('原生八方向生产门仍关闭。已保存的修补版全 1024 盒首个 h=0.005 plant 条件性门覆盖 20,480 个合成物理态界；lane 0 第二步门只涉及一盒。10 月 4 日控制余项构造回放已检查 1024 盒、3072 输出、196608 个精确仿射顶点，四个阶段 exit 0，但依赖原实数 CROWN 包络有效，未调用 NN/CROWN 或 ODE。它不认证后续控制或整个时域。'+link('已完成回放与限制','docs/evidence/results/archcomp26_20261001/native_quad_allbox_adaptive_remainder_replay_20261004_001/README.md')+'。')
        if key=='single-pendulum-reach':add('x3 空栏专门保留固定 MATLAB 第三返回量的来源缺口；论文的两个物理态结果仅为具名两态 profile，不能凭 dx3=1 自行补第三初值和执行入口。')
        if key=='tora-remain':add('另有 h=0.05 四方全程补充合同，4800/4800 盒步均保存；T=20 的 x4 endpoint 并集宽度 P3/H/X/原生为 0.354314/0.408728/0.408728/0.346793。该步长变体不替换 h=0.1 主格。'+link('独立合同完整四态宽度与时间来源','docs/evidence/results/archcomp26_20261001/tora_remain_h005_fourway_saved_20261002/SUMMARY.md')+'。')
        if key=='airplane-continuous':add('64 个二分子盒各完成首个 0.01 秒 plant 步，其中 8 盒保存安全、56 盒 Unknown；高角子盒随后只完成 4/10 小步到 0.04 秒，第 5 步 x/y 自包含失败。这是另列数值诊断，不填主合同的 T=2 格，也不是实际不安全轨迹。')
        if key in ('tora-reach-sigmoid','tora-reach-tanh'):
            tag='sigmoid' if key.endswith('sigmoid') else 'tanh'
            add(f'![{tag} 四态逐步 endpoint 宽度](evidence/results/archcomp26_report_20261006/figures/tora_{tag}_endpoint_widths.png)')
            add(f'![{tag} 四态逐步 tube 宽度](evidence/results/archcomp26_report_20261006/figures/tora_{tag}_tube_widths.png)')
        source_ids=sorted({sid for m in METHODS for s in states for sid in ws[key,m,s]['source_ids'].split(';') if sid})
        add('宽度来源编号：'+(', '.join(source_ids) if source_ids else '无有效主合同范围')+'；'+link('来源路径与大小','docs/evidence/results/archcomp26_report_20261006/widths/sources.json')+'、'+link('上下界和全部逐步宽度','docs/evidence/results/archcomp26_report_20261006/widths/widths_long.csv')+'、'+link('缺失逐格说明','docs/evidence/results/archcomp26_report_20261006/widths/missing_fields.csv')+'。')
    add('## 附录 A 本轮候选和未采用原因')
    add('下表保留本轮全部独立尝试。小步完成数与包装资格结果分开：数值算完不代表通过实现门；短试验不外推为完整时域。下表不是不同合同的速度排行榜。')
    rows=[]
    for r in runs:
        label=r['run_id'].removesuffix('_001').replace('quad_paper_','QUAD ').replace('tora_sigmoid_','TORA sigmoid ').replace('tora_tanh_','TORA tanh ').replace('nav_standard_','NAV standard ').replace('nav_robust_','NAV robust ').replace('sp_two_state_','SP ').replace('attitude_','Attitude ').replace('docking_','Docking ').replace('_full',' ').replace('_',' ')
        full=r.get('horizon',{}).get('complete_named_horizon')
        valid=r.get('outer_exit_code')==0 and str(r.get('status','')).startswith('COMPLETED')
        state=('完整' if full else '短测通过') if valid else ('资格未过' if r.get('completed_substeps',0) else '首步前错误')
        rows.append([link(label,r['raw_result']),str(r.get('completed_substeps') if r.get('completed_substeps') is not None else '无已记录 ODE'),sec(r.get('driver_elapsed_s')),sec(r.get('outer_wall_s')),state])
    add(table(['候选及原始 RESULT','完成小步','driver','process','状态'],rows))
    notes=STUDY/'REPORT_CANDIDATE_NOTES.md'
    if notes.is_file(): add(notes.read_text().strip())
    add('10 月 5 日 QUAD two-slope 短试验也保留：40 步中 x5 较窄，但 x9 在大多数步更宽，未作为默认。完整历史候选见 '+link('上一版速度研究','research/p3_speed_tightness_20261005/README.md')+'，没有重跑。')
    add('## 附录 B 图形功能与尚缺证据')
    add('Python 绘图 CLI 继续支持现有 ranges.bin 和兼容几何、初盒、tube / endpoint、按时间定义的 Safe / Target、同轴多方法 PNG / PDF、JSON / CSV。实现提速保留原逐步输出，新 TORA / QUAD 图使用各自新候选范围。当前绘图入口不生成 MATLAB 文件；已有归档保持原样，Python 图、数据和原有输入功能均保留。')
    add('独立 NNCS 浮点端到端证明和原生八方向生产门仍未闭合。TORA endpoint 落入所选目标只是保存几何观察，原 GPU property_evaluated=false 不会变成正式 reach 证书。Docking 完整数值但 UNKNOWN；SP 官方第三态、Balancing 五特征模型、Airplane discrete 权威转换合同仍缺件。其他数值拒绝原因和完整前缀均在对应分节逐方法说明。')
    add('## 附录 C 当前报告与原始来源')
    add('当前正文与 10 月 6 日报告包同步；'+link('上一版正文快照','docs/ARCHCOMP26_REPORT_HISTORY_20261005.md')+'和旧日期 Word / PDF 作为历史保留。原 289 条尝试、10 月 5 日 13 个优化资格阶段、本轮新候选分别列账。报告重建只读取已保存证据，不启动求解器、旧检查器或摘要计算。')
    add('- '+link('Word PDF 和全部数据入口','docs/evidence/results/archcomp26_report_20261006/README.md'))
    add('- '+link('64 格具体状态与未完成原因','docs/evidence/results/archcomp26_report_20261006/blockers/README.zh.md'))
    add('- '+link('逐进程全部时间层','docs/evidence/results/archcomp26_report_20261006/timing/runs.csv')+'；'+link('完整时间索引','docs/evidence/results/archcomp26_report_20261006/timing/timing_index.json'))
    add('- '+link('逐态宽度概要','docs/evidence/results/archcomp26_report_20261006/widths/summary.csv')+'；'+link('全部逐步范围','docs/evidence/results/archcomp26_report_20261006/widths/widths_long.csv')+'；'+link('逐态差值与包含','docs/evidence/results/archcomp26_report_20261006/widths/pairwise_comparisons.csv'))
    SOURCE.write_text('\n\n'.join(lines)+'\n')
    (OUT/'REPORT_ASSEMBLY.json').write_text(json.dumps(dict(instance_sections=16,status_cells=len(cells),width_summary_rows=len(width),pairwise_rows=len(comparisons),new_stages=len(runs),solver_runs=0,old_checker_runs=0,digest_operations=0,report_source=str(SOURCE.relative_to(ROOT))),indent=2)+'\n')
    print(f'Markdown: {len(lines)} blocks; {len(comparisons)} pairwise comparisons')

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--docx',action='store_true');parser.add_argument('--render-dir',type=Path);args=parser.parse_args()
    assemble()
    if args.docx or args.render_dir:
        helper=docx_build()
        if args.render_dir:helper.render(args.render_dir)
