"""Build the editable ARCH-COMP26 progress report from saved, no-hash evidence."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "docs/evidence"
OUT = EVIDENCE / "results/archcomp26_20261001/stage_report/ARCHCOMP26_STATUS_20261002.docx"
FONT = "Hiragino Sans GB"
METHODS = ("pytorch_gpu", "huan", "xiangru", "flowstar_native")
LABELS = {
    "current_numeric_full": "新全程",
    "historical_same_contract_numeric_full": "旧全程",
    "no_full_numeric_evidence": "无全程",
    "contract_blocked": "合同未定",
}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def font_for_cjk(font, element):
    font.name = FONT
    rpr = element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    edge = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        line = OxmlElement(f"w:{side}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), "D9D9D9")
        edge.append(line)
    tc_pr.append(edge)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    for index, row in enumerate((headers, *rows)):
        cells = t.rows[0].cells if index == 0 else t.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            cell.text = str(value)
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, "DFE8F1" if index == 0 else ("F7F9FB" if index % 2 == 0 else "FFFFFF"))
            borders(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    font_for_cjk(run.font, run._element)
                    run.font.size = Pt(8.0 if len(rows) > 10 else 8.5)
                    run.bold = index == 0
        if index:
            t.rows[index]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph()
    return t


def paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def main():
    overlay = json.loads((EVIDENCE / "archcomp26_coverage_overlay_20261002.json").read_text())
    matrix = json.loads((EVIDENCE / "archcomp26_nohash_work_matrix_20261001.json").read_text())
    cells = {(x["instance_id"], x["method"]): x for x in overlay["cells"]}
    assert len(cells) == matrix["cell_count"] == 64
    assert overlay["new_attempt_count"] == matrix["attempt_count"]
    assert sum(overlay["coverage_counts"].values()) == 64
    names = list(dict.fromkeys(x["instance_id"] for x in overlay["cells"]))
    assert len(names) == 16 and all((name, method) in cells for name in names for method in METHODS)

    doc = Document()
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(1.75)
        section.footer.paragraphs[0].text = "ARCH-COMP26 非 VCAS 证据进度  ·  2026-10-02"
    normal = doc.styles["Normal"]
    font_for_cjk(normal.font, normal._element)
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 17), ("Heading 1", 12), ("Heading 2", 10)):
        style = doc.styles[name]
        font_for_cjk(style.font, style._element)
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    for edge in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(edge)

    doc.add_paragraph("ARCH COMP26 非 VCAS 四方实验进度", "Title")
    paragraph(doc, "2026 年 10 月 2 日 · 可审阅阶段稿；实验和最终报告尚未完成。")
    counts = overlay["coverage_counts"]
    paragraph(doc, (
        f"已登记 {matrix['attempt_count']} 条本轮尝试，覆盖 16 个实例 × 4 种方法。"
        f"64 格中，本轮完整数值时域 {counts['current_numeric_full']} 格；"
        f"经同合同审计可复用的旧全程 {counts['historical_same_contract_numeric_full']} 格；"
        f"无完整数值时域 {counts['no_full_numeric_evidence']} 格；"
        f"Airplane 离散语义待定 {counts['contract_blocked']} 格。"
        "完整数值时域、作者性质标签、独立端到端浮点 NNCS 证明与稳定速度排名是不同口径。"
    ))

    doc.add_heading("16 个实例的数值时域覆盖", 1)
    table(doc, ("实例", "P3 GPU", "Huan", "Xiangru", "Flow*"), [
        (name, *(LABELS[cells[name, method]["coverage_status"]] for method in METHODS))
        for name in names
    ], (6.6, 2.6, 2.6, 2.6, 2.6))
    paragraph(doc, "新全程只说明本轮保存了全部数值时间段；旧全程仍留在原历史记录中，未计入 187 条新尝试。无全程包含真实首步失败或后期数值拒绝。任何格均不自动获得性质证明或计时排名资格。")

    doc.add_heading("已得到的关键结果", 1)
    for title, body in (
        ("Huan QUAD 约 80 秒", "历史作者合同下，parity 加保存盒、原生 float64 与 SR 分块的 1024 盒 × 1000 步五次进程中位数为 75.250099 秒；strict 同盒变体在第 597 步拒绝。P3 加 sin/cos 复用的完整单次为 1533.752052 秒，较其旧 P3 单次约快 2.056 倍。两种合同及保证边界在专项报告中分开。"),
        ("2026 论文方程 QUAD", "四方均完成 1024 盒 × 1000 步；T=5 时 x3 终点并集位于目标 [0.94,1.06]。P3、Huan/Xiangru、原生的 x3 并集宽分别为 0.067225749、0.047741817、0.050976637。原生 checker 的 VERIFIED 是保存入口的终点检查，仍缺论文 reach and remain 的权威全时间窗执行语义。"),
        ("Unicycle 与 TORA reach sigmoid", "按用户指定的论文常值速度扰动，Unicycle 四方均完成 T=10；P3 与原生的保存终点全盒入目标，Huan/Xiangru 终点未全入，性质仍为 Unknown。TORA 官方 sigmoid、u=11f 四方均完成 T=5 且数值终点入目标；只有原生执行作者终点 checker。两项均无独立端到端浮点证明。"),
        ("三个具名计时组合", "ACC participant order、Single Pendulum 两物理态、修正危险集 Attitude 各有四方首轮加五次后续独立进程。其合同、观察器和验证资格不同；报告逐次给出原始时间，不据此跨方法排稳定名次。"),
    ):
        doc.add_heading(title, 2)
        paragraph(doc, body)

    doc.add_heading("仍需解决的具体问题", 1)
    for item in (
        "Airplane 离散：缺参与者权威离散转移与控制更新顺序；四方法主表暂不运行猜定合同。",
        "Single Pendulum：缺官方三态第三状态初值与 MATLAB 闭环执行入口；现有四方是两物理态加辅助时钟合同。",
        "Balancing：缺论文五输入控制器或到官方四输入模型的明确映射；固定仓库 raw4 四方均早停。",
        "QUAD reach and remain：缺参与者全时间窗 checker 源码或等价权威执行记录。",
        "Airplane 连续、DP more、TORA remain 作者两方均有数值拒绝或资源阻断；TORA remain 的 Unknown 性质标签并不是数值拒绝原因，关闭性质检查不能补齐流管。",
    ):
        paragraph(doc, item, "List Bullet")

    doc.add_heading("14 格无完整数值时域的实测停止点", 1)
    table(doc, ("实例", "格数", "已保存的停止点"), (
        ("Airplane continuous", "4", "全初盒四方首步均未得到可用流管；Huan 高阶另有资源阻断。"),
        ("Balancing raw4", "4", "P3 第 87 步；Huan/Xiangru 第 99 步；原生保存 83 段后停止。论文五特征合同另缺控制器。"),
        ("DP more robust", "4", "P3 第 9 步数值拒绝；Huan/Xiangru 各保存 72/80 步，原生 64/80 步。"),
        ("TORA remain", "2", "Huan/Xiangru 各 2357/2400 盒步接受；第 190 步首次盒拒绝，最终各仅 6/12 盒接受。"),
    ), (4.6, 1.2, 11.2))
    paragraph(doc, "这些停止点不外推为完整进程时间。早停的保存区间跨安全边界是性质 Unknown 或不能证实；不能直接称为实际不安全轨迹。")

    doc.add_heading("证据入口", 1)
    for path in (
        "docs/ARCHCOMP26_FINAL_REPORT_DRAFT.md",
        "docs/evidence/archcomp26_nohash_attempts_20261001.json",
        "docs/evidence/archcomp26_coverage_overlay_20261002.md",
        "docs/ARCHCOMP26_TORA_REMAIN_NUMERIC_STOP_AUDIT_20261002.md",
        "docs/HUAN_QUAD_SPEED_AND_MODES.md",
    ):
        paragraph(doc, path)
    paragraph(doc, "所有路径均相对于仓库根目录；原始 RESULT、日志和已保存数值范围由索引指向。报告与覆盖表不运行求解器，也不执行内容摘要校验。")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
