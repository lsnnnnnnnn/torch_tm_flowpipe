"""Rebuild the editable stage DOCX from its pre-ACC/Attitude copy.

Reads saved ARCH-COMP26 evidence only. The PDF is exported from the resulting
DOCX in Microsoft Word so Chinese typography survives. No solver is launched.
"""
from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "docs/evidence/results/archcomp26_20261001"
STAGE = RESULTS / "stage_report"
BASE = STAGE / "ARCHCOMP26_STAGE_REPORT_DRAFT_PRE_ACC_ATTITUDE_20261001.docx"
OUTPUT = STAGE / "ARCHCOMP26_STAGE_REPORT_DRAFT_20261001.docx"
PLOTS = RESULTS / "plots/nohash_saved_20261001"
WIDTHS = RESULTS / "acc_fourway_saved_ranges_20261001/acc_t5_endpoint_and_full_tube_wide.csv"


def _find(doc: Document, prefix: str):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"expected one paragraph starting {prefix!r}, found {len(matches)}")
    return matches[0]


def _insert(anchor, text: str = "", style: str = "Normal", *, page_break: bool = False):
    paragraph = anchor.insert_paragraph_before(text, style=style)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:color"), "D9D9D9")
        edge.set(qn("w:sz"), "5")
        borders.append(edge)


def _pad(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:w"), "95")
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    tc_pr.append(margins)


def _width_table(doc, anchor):
    with WIDTHS.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    if len(rows) != 6 or [r["state"] for r in rows] != [
        "x_lead", "v_lead", "a_lead", "x_ego", "v_ego", "a_ego"
    ]:
        raise ValueError("ACC saved endpoint CSV must contain the six physical states")
    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    widths = [Cm(3.1), Cm(3.55), Cm(3.55), Cm(3.55), Cm(3.55)]
    headers = ("物理态", "native", "Huan", "Xiangru", "ours/P3")
    for i, title in enumerate(headers):
        table.rows[0].cells[i].text = title
    keys = ("native_endpoint_t5_width", "huan_endpoint_t5_width",
            "xiangru_endpoint_t5_width", "ours_p3_endpoint_t5_width")
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row["state"]
        for index, key in enumerate(keys, 1):
            cells[index].text = f"{float(row[key]):.8f}"
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _border(cell)
            _pad(cell)
            if row_index == 0:
                _shade(cell, "DDEAF4")
            elif row_index % 2 == 0:
                _shade(cell, "F6F9FC")
            for paragraph in cell.paragraphs:
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if column_index == 0
                                       else WD_ALIGN_PARAGRAPH.CENTER)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.4)
                    if row_index == 0:
                        run.bold = True
    anchor._p.addprevious(table._tbl)
    return table


def _campaign_table(doc, anchor):
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    widths = [Cm(4.3), Cm(3.2), Cm(3.3), Cm(5.2)]
    values = [
        ("方法", "首次新进程 (s)", "后 5 次中位数 (s)", "后 5 次 min–max (s)"),
        ("Flow* native RPC", "7.686284", "7.736304", "7.636556–7.786681"),
        ("Huan GPU", "8.087964", "8.237208", "8.138279–8.337823"),
        ("Xiangru GPU", "7.887486", "7.987532", "7.937743–8.086614"),
        ("ours/P3", "8.589842", "8.740510", "8.637965–8.840520"),
    ]
    for row_index, data in enumerate(values):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for column_index, (cell, value) in enumerate(zip(cells, data)):
            cell.text = value
            cell.width = widths[column_index]
            _border(cell)
            _pad(cell)
            if row_index == 0:
                _shade(cell, "DDEAF4")
            elif row_index % 2 == 0:
                _shade(cell, "F6F9FC")
            for paragraph in cell.paragraphs:
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if column_index == 0
                                       else WD_ALIGN_PARAGRAPH.CENTER)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.4)
                    if row_index == 0:
                        run.bold = True
    anchor._p.addprevious(table._tbl)
    return table


def _figure(anchor, path: Path, caption: str, note: str) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    image = _insert(anchor)
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.keep_with_next = True
    image.add_run().add_picture(str(path), width=Cm(16.8))
    cap = _insert(anchor, caption, "Caption")
    cap.paragraph_format.keep_with_next = True
    _insert(anchor, note, "Note")


def main() -> None:
    doc = Document(BASE)
    _find(doc, "2026 年 10 月 1 日").text = "2026 年 10 月 2 日  ·  阶段更新版"
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                if "2026-10-01" in run.text:
                    run.text = run.text.replace("2026-10-01", "2026-10-02")
    updates = {
        "目标是在同一 benchmark 合同下": (
            "目标是在同一 benchmark 合同下，比较 PyTorch/GPU、Huan、Xiangru 和 Flow* native 对 16 个非 VCAS 实例的完整性、全进程时间及绝对 flowpipe 宽度；计划矩阵为 16×4，共 64 个 cell。当前有 24 个方法单元的完整数值时域记录，其中 ACC participant-order、两物理态 Single Pendulum、修正危险盒的 Attitude 与 Docking 各有四方完整时域记录；Docking 性质未决，四方矩阵仍未完成。"
        ),
        "Single Pendulum 两物理态": (
            "两物理态 Single Pendulum 的四方各完成 1×100 子步，P3 的 50 个性质窗口安全事件与保存 tube 均在带内；这不是官方三物理态 MATLAB 复现。TORA remain 原生与 P3 各完成 12×200，Huan/Xiangru 仅有 t≤18.4 的共同全盒接受且保存 tube 安全前缀。DP more 三方未完成 T=0.4，P3 的 DP less 未完成 225 盒全程。ACC 四方各完成完整初盒的 50 控制期，并额外各有 1+5 次独立新进程计时；修正官方 unsafe 盒后的 Attitude 四方各完成 30 期、60 小段。Docking 四方各完成 40 期、400 小段，但全时非线性安全性质均为 Unknown；Airplane continuous 全初盒四方入口均未得到首个接受小步。保存区间、作者 checker 输出和图均不构成独立端到端浮点 NNCS 证明。"
        ),
        "正式时间目标为": (
            "正式时间目标为 1 次冷启动及 5 次独立进程 steady。ACC participant-order 四方另有各 1 次首轮和 5 次后续独立新进程；其余完整方法单元只有单次样本。该共享主机 campaign 的中位数仅作描述，不给稳定速度排名。各方法数值引擎、观察器和冷启动构成不同；TORA Huan/Xiangru 的失败进程耗时不能与完整时间排名。"
        ),
        "论文方程 QUAD：Flow* native": (
            "论文方程 QUAD：Flow* native 全程尚在运行；P3/Huan/Xiangru 的重复计时待做。P3 的已保存 pooled union t–x₃ tube 已绘图，但它不是逐盒投影。"
        ),
        "Single Pendulum 与 TORA remain：": (
            "两物理态 Single Pendulum 四方已有单次全程，但与官方三物理态 MATLAB 合同不同；TORA 原生/P3 完整与 Huan/Xiangru 晚期拒绝的性质语义须并列报告。Attitude 修正 unsafe 后四方单次全程均已记录；Docking 四方也完成数值时域但性质 Unknown。Airplane continuous 的完整初盒四方入口遇到资源、编码容量或首小步拒绝。其他实例仍须逐项冻结合同并记录尝试。"
        ),
        "当前保存区间、图和作者 checker": (
            "ACC 在明确命名的 participant-order 合同下已有四方单次完整数值记录，但没有稳定四方速度排名。保存区间、图和作者 checker verdict 不是独立端到端浮点证明。TORA Huan/Xiangru 的区间出带不是独立实际轨迹反例；失败进程时间也不是 T=20 完成时间。路径与配置只作来源说明，未与记录内容绑定；旧冻结证据仍仅为历史参照。"
        ),
        "“全程”表示本轮一个方法": (
            "“全程”表示本轮一个方法完成目标数值时域，并不自动表示性质已证明；“诊断”表示短程入口尝试；“早停”或“未知”表示未完成。ACC、两物理态 Single Pendulum、Attitude 修正后、Docking 各四方及 TORA 原生/P3 等共 24 格有完整记录，距离 64 格矩阵仍有缺口。Docking 原生的 outer RESULT 是 failed/exit 2，原因是 checker Unknown，数值仍完成 400 段。ACC 另有 24 次独立新进程 campaign；其计时中位数不构成稳定速度排名。TORA Huan/Xiangru 虽记录至第 200 小步，但部分盒步拒绝，故列为“未知·前缀”。"
        ),
    }
    for old, new in updates.items():
        _find(doc, old).text = new

    # The historical baseline caveat is stated in the abstract and appendix;
    # the standalone sentence otherwise occupies an entire Word page.
    baseline_note = _find(doc, "旧基线矩阵中的 not_started=64")
    baseline_note._element.getparent().remove(baseline_note._element)
    repeated_qualification = _find(doc, "四方共同前缀与排名资格需要")
    repeated_qualification._element.getparent().remove(repeated_qualification._element)
    dp_heading = _find(doc, "2  Double Pendulum less-robust")
    preceding = dp_heading._p.getprevious()
    if preceding is not None and 'w:type="page"' in preceding.xml:
        preceding.getparent().remove(preceding)

    _find(doc, "9  Single Pendulum · 新三方全程").text = "9  Single Pendulum · 两物理态四方全程"
    for title in ("8  P3 诊断、旧证据与后续工作", "9  Single Pendulum · 两物理态四方全程"):
        preceding = _find(doc, title)._p.getprevious()
        if preceding is not None and 'w:type="page"' in preceding.xml:
            preceding.getparent().remove(preceding)
    _find(doc, "合同是两物理态加辅助时钟").text = (
        "合同是两物理态加辅助时钟：一个初盒、20 控制期、100 个小步；性质只在 t∈[0.5,1] 要求 x₁∈[0,1]，不能将该带延伸到全时域。原生、Huan、Xiangru、P3 各有单次完整保存。P3 接受 100/100 小步，窗口 50 个安全事件均通过；独立扫描这 50 个保存 tube 的 x₁ union 为 [0.5645452654370386,0.9932406822927875]，位于指定带内。该合同并非官方三物理态 MATLAB 复现。"
    )
    _find(doc, "进程 wall 分别为原生").text = (
        "单次外层进程 wall 分别为原生 4.726027 s、Huan 5.430762 s、Xiangru 5.528731 s、P3 6.583305 s。数值引擎与资源配置不同，单次时间及保存区间不能形成稳定四方排名或独立端到端浮点证明。详见 single_pendulum_prep_001/SUMMARY.md、native_sp_two_state_full20_001/SUMMARY.md 与 single_pendulum_two_state_p3_full20_001/SUMMARY.md。"
    )
    _find(doc, "10  TORA remain · 原生全程与作者前缀").text = "10  TORA remain · 原生/P3 全程与作者前缀"
    _find(doc, "固定合同：初盒").text = (
        "固定合同：初盒 [0.6,0.7]×[-0.7,-0.6]×[-0.4,-0.3]×[0.5,0.6] 分成 12 盒；T=20，控制周期 1，200 个小步；四态全时安全带 [-2,2]。原生保存 2400/2400 条范围，作者 checker 报 VERIFIED，全部保存 tube 在安全带内；单次进程 wall 8.339151 s。P3 新全程 2400/2400 盒步 accepted，独立扫描 2400 条保存范围确认全部 tube 在安全带内，单次外层 wall 10.411001 s；P3 日志未给出可等同原生 VERIFIED 的显式终判。"
    )

    evidence = doc.tables[0]
    for row in evidence.rows:
        if row.cells[0].text.startswith("Single Pendulum · native/Huan/Xiangru"):
            row.cells[0].text = "Single Pendulum 两物理态 · 四方 · 新 2026"
            row.cells[1].text = "各 1 盒 × 100 子步；20 控制期"
            row.cells[2].text = "四方单次全程；P3 100/100 accepted，窗口 50 个安全事件通过，保存 tube 在目标带内；非官方三物理态 MATLAB 复现。"
        if row.cells[0].text.startswith("TORA remain · native 与 Huan/Xiangru"):
            row.cells[0].text = "TORA remain · native/P3 与 Huan/Xiangru · 新 2026"
            row.cells[2].text = "原生/P3 全程且保存 tube 安全；原生显式 VERIFIED，P3 2400/2400 accepted；两方仅接受 2357/2400，共同合格前缀 t≤18.4，作者 Unknown."
    insert_at = next(i for i, row in enumerate(evidence.rows)
                     if row.cells[0].text.startswith("DP less · P3"))
    for cells in (
        ("ACC participant-order · 四方 · 新 2026", "各 1 盒 × 50 控制期；T=5，全时半空间",
         "四方单次全程，保存 tube 的 margin 下界均正；尚无稳态速度排名或端到端证书。"),
        ("Attitude avoid · 修正 unsafe · 四方新 2026", "各 1 盒 × 30 期 × 2 小段；T=3",
         "原生显式 VERIFIED；Huan/Xiangru checker 未打印 Unsafe/Unknown；P3 60/60 accepted；四方各 60 个保存 tube 盒与官方 unsafe 盒不相交。"),
        ("Docking · 完整初盒 · 四方新 2026", "各 1 盒 × 40 期 × 10 小段；T=40",
         "四方各完成 400 段数值时域，非线性全时性质均 Unknown；原生 outer exit 2 如实保留，未将区间相交当实际反例。"),
        ("Airplane continuous · 完整初盒入口", "12 物理态单盒；目标 20 期，T=2",
         "Huan order 6 资源阻断；Huan/Xiangru order 3、P3 严格同阶及 native 均首步未接受；P3 另一验证设置在建表时溢出。无完整数值结果。"),
    ):
        row = evidence.add_row()
        for cell, value in zip(row.cells, cells):
            cell.text = value
            _border(cell)
            _pad(cell)
            if cells[0].startswith("ACC"):
                _shade(cell, "F6F9FC")
        evidence._tbl.remove(row._tr)
        evidence._tbl.insert(insert_at + 1, row._tr)
        insert_at += 1
    for row in evidence.rows:
        if row.cells[0].text.startswith(("Single Pendulum 两物理态", "TORA remain")):
            for cell in row.cells:
                _border(cell)
                _pad(cell)
                _shade(cell, "F6F9FC")
    coverage = doc.tables[5]
    for row in coverage.rows[1:]:
        name = row.cells[0].text
        if name == "single-pendulum-reach":
            row.cells[1].text = "全程"
        if name == "acc-safe-distance":
            for index in range(1, 5):
                row.cells[index].text = "全程"
        if name == "attitude-control-avoid":
            row.cells[1].text = "全程"
            row.cells[2].text = "全程"
            row.cells[3].text = "全程"
            row.cells[4].text = "全程"
        if name == "tora-remain":
            row.cells[1].text = "全程"
        if name == "docking-constraint":
            for index in range(1, 5):
                row.cells[index].text = "全程·性质未知"
        if name == "airplane-continuous":
            for index in range(1, 5):
                row.cells[index].text = "失败"

    # Keep the pooled QUAD drawing with its existing contract section.
    anchor = _find(doc, "7  DP more-robust")
    _insert(anchor, "6.1  P3 QUAD 已保存的 t–x₃ pooled tube", "Heading 2", page_break=True)
    _insert(anchor, "每个局部小步只有 1024 个已接受盒的 x₃ 合并上下界，没有逐 lane 二维几何。图示为 1000 个保存 tube 的无插值包络；[0.94,1.06] 只在 T=5 标为终点目标。")
    _figure(anchor, PLOTS / "quad_paper_p3_1024x1000_t_x3_pooled_tube.png",
            "图 3  新论文方程 QUAD P3，1024 盒 × 1000 步的 pooled t–x₃ tube；不是逐盒投影。",
            "最后一行 observer 的 x₃ endpoint 与驱动最终 hull 数值略异，图忠实采用保存 JSONL。见 quad_paper_p3_nohash_v1/full50_001/data/observations.jsonl；MATLAB 脚本未实跑。")

    # New full-coverage example and the corrected unsafe-set result.
    anchor = _find(doc, "附录 A")
    _insert(anchor, "11  ACC · participant-order 四方全程", "Heading 1", page_break=True)
    _insert(anchor, "共同命名合同采用完整初盒 [90,110]×[32,32.2]×{0}×[10,11]×[30,30.2]×{0}，一个 lane，T=5、50 次 0.1 s 控制保持。控制器输入第五项固定为 v_lead−v_ego；2026 论文未定义这个符号，故本表仅属于该 participant-order profile。全时性质为 x_lead−x_ego−1.4v_ego−10≥0，不是轴对齐 Safe 盒。")
    _insert(anchor, "native corrected-VAR、Huan、Xiangru、ours/P3 四方均有单次完整 50 期保存与接受/运行证据；各自进程 wall 为 7.930、7.260、7.312、7.934 s。下面列出第 50 期保存 endpoint 的六态区间宽度，单位沿用物理态原单位；绝对上下界及全时 tube union 见 acc_fourway_saved_ranges_20261001 的 CSV。观察器不同、冷启动路径不同，不据单次时间或微小宽度差异作速度和精度排名。")
    _insert(anchor, "表 6  ACC 四方法 T=5 保存 endpoint 的六态绝对区间宽度（显示至 8 位小数）", "Caption")
    _width_table(doc, anchor)
    _insert(anchor, "四方的 50 段保存 tube 逐段按完整六态盒计算保守 safe-distance margin，最低下界依次为 native 16.254753759211283、Huan 16.165034707564207、Xiangru 16.165034707564207、ours/P3 16.434858569716976，全部大于零。系数 1.4 作为精确 7/5，保存 binary64 盒端点的有理数像在输出时向外取整；坐标相关性不可恢复。native ranges.bin 不含 accepted/status；GPU 三方保存行均为 accepted=true。")
    _insert(anchor, "11.1  ACC 四方独立新进程计时 campaign", "Heading 2")
    _insert(anchor, "同名 participant-order 合同另有 24 次新进程：四方各一条首次进程及五条后续独立进程，24/24 退出正常并完成 50 期；每次保存 tube 的安全距离下界都大于零。均在同一物理 GPU 2、CPU 10–13 上顺序运行，但 GPU 1 同时有其他作业。首次进程也不是重启主机后的冷机；supervisor 从子进程启动前计至回收后，小于约 0.05 s 的差异不作性能解释。下表仅为描述性计时，不构成稳定速度排名；24 次明细见 acc_fourway_campaign_001/SUMMARY.md、RUNS.csv 与 INDEPENDENT_AUDIT.json。")
    _insert(anchor, "表 7  ACC campaign：四方法各 1 次首次进程与后 5 次独立进程 wall", "Caption")
    _campaign_table(doc, anchor)
    _insert(anchor, "11.2  ACC 四方保存 tube 的 t–margin 图", "Heading 2", page_break=True)
    _figure(anchor, PLOTS / "acc_four_method_t_safe_distance_margin_tube.png",
            "图 4  ACC 四方每期 whole-tube 盒的保守 safe-distance margin；虚线为全时 margin=0 边界。",
            "图源为各方 50 期保存范围，不是 CSV 的全时 union，也不是独立浮点 NN 证明。图源路径、大小和方法配置见 plots/nohash_saved_20261001 的 geometry 与 render 收据；MATLAB 脚本未实跑。")

    _insert(anchor, "12  Attitude Control · 修正危险盒后的四方全程", "Heading 1", page_break=True)
    _insert(anchor, "固定官方 unsafe 盒的第 4 坐标约束是 −0.7≤x₄≤−0.6；旧 native/Xiangru checker 误写成 x₄≥−0.4 且 x₄≤−0.6，形成空集，所以旧 VERIFIED 不能证明官方性质。新隔离构建修正该条件后，原生、Huan、Xiangru、P3 用完整初盒各完成 30/30 控制期与 60/60 本地小段。单次外层进程 wall 分别为 6.281498、7.094217、6.933376、13.017036 s；P3 60/60 accepted。只有原生 checker 显式打印 VERIFIED，Huan/Xiangru 的 checker 未打印 Unsafe/Unknown。")
    _insert(anchor, "独立扫描四方新保存的各 60 个六态 tube 盒，均与官方 unsafe 盒不相交；原生第 4 坐标全时 tube 上界 −0.710800050132425，低于 unsafe 下界 −0.7。该盒分离结论只涉及保存的轴对齐包络；上游神经网络边界与所有浮点步骤尚无独立端到端证明。不能沿用旧错误 checker 的成绩，也不能据单次时间作速度排名。来源见 native_attitude_avoid_full30_001/SCAN.json、BOX_UNSAFE_AUDIT.json、author_attitude_avoid_v1、p3_attitude_avoid_v1/full30_001 与 ARCHCOMP26_ATTITUDE_CONTROL_CONTRACT_20261001.md。")

    _insert(anchor, "13  Docking · 完整初盒四方数值时域", "Heading 1")
    _insert(anchor, "固定 2026 合同以一盒 [70,106]²×[−0.28,0.28]² 覆盖四态 (sx,sy,vx,vy)，控制器原始四态输入、两力输出，1 s 保持并更新 40 期。四方各记录 400 个 0.1 s 全时 tube；外层进程 wall 为 P3 17.411771 s、Huan 12.698839 s、Xiangru 12.668129 s、原生 9.139416 s，均只有一个完整进程样本。原生 wrapper 的原始状态是 failed/exit 2，原因是全时性质 UNKNOWN；数值日志独立显示 40/40 期、400 tubes、40 RPC，不可把外层状态改写成 VERIFIED。")
    _insert(anchor, "全时安全性质为 q=√(vx²+vy²)−0.2−0.002054√(sx²+sy²)≤0。初盒 q 上界约 −0.007356；首段 [0,0.1] 的四方保守 q 上界均约 +0.016083。四方的 400 段保存 tube 盒检查均不能证明 q≤0；盒与 q>0 侧相交不构成实际轨迹反例。原生末段最大 q 上界 +8.954347，其余三方约 +8.806 至 +8.808。四方末时绝对区间、原生上下仿射斜率一致性与原始 RESULT 见 DOCKING_FULLBOX_3METHODS_SUMMARY.md、native_docking_full40_001/SUMMARY.md。")
    _figure(anchor, PLOTS / "docking_fullbox_4method_q_upper.png",
            "图 5  Docking 四方保存 tube 的 q 上界；绿色为 q≤0 安全区域，内框放大第一秒。",
            "每条曲线来自完整初盒保存的 400 个区间，不是轨迹或安全证明。原生外层 exit 2 与四方 Unknown 均保留；图源与 MATLAB 脚本在 plots/nohash_saved_20261001，MATLAB 未实跑。")

    _insert(anchor, "14  Airplane continuous · 完整初盒四方入口阻断", "Heading 1")
    _insert(anchor, "固定官方连续版把 (u,v,w,phi,theta,psi) 六维设为 [0,1]，其余六物理态为零，且只给一个未分割初盒；控制器 12 输入、6 输出，20 个 0.1 s 周期至 T=2。全时要求 y、phi、theta、psi∈[−1,1]。旧点初盒的四方成绩不覆盖该盒。Huan 的 order 6 一期尝试在零 ODE 小步前因 19 变量六阶单项式配对表导致 RSS 超过 54,006,540 KiB 而停；独立 order 3 诊断中 Huan/Xiangru 均在首个 0.01 s 小步拒绝唯一全初盒，接受数为零。三次外层 wall 分别为 252.247、6.785、6.985 s，均非完整 T=2 性能样本。")
    _insert(anchor, "P3 严格 solution_plus_one 首次新 smoke 在建验证 P4 表时发生 64 位整数编码溢出，0 段接受，wall 13.757866 s；另一新 run 使用引擎现有的严格同阶 solution_order 验证，首个 0.01 s 小步 accepted=false，0/10 段接受，wall 12.500084 s。Flow* native 独立 order 6 和 order 3 smoke 均在首个 0.01 s 小步返回 status 4=UNCOMPLETED_SAFE，0 flowpipes，外层 failed/exit 2；wall 分别为 4.877421 和 3.925546 s。order 3 把先验 remainder 从 [−0.01,0.01] 扩至 [−1,1] 的额外诊断仍在首步返回 status 4，这个改参结果不属于官方基线。均没有全时 tube 或性质结论，失败时间不能作完整性能比较；P3 内部拒绝细因未记录。原始摘要见 airplane_continuous_order3_fullbox_20261002/SUMMARY.md、AIRPLANE_P3_FULLBOX_SMOKES_20261002.md 与 native_airplane_fullbox_smokes_20261002/SUMMARY.md。")

    # Small provenance index without repeating large source tables.
    anchor = _find(doc, "DP 绘图源与输出")
    for path in (
        "results/archcomp26_20261001/ACC_PARTICIPANT_PROFILE_20261001.md",
        "results/archcomp26_20261001/acc_fourway_campaign_001/SUMMARY.md",
        "results/archcomp26_20261001/acc_fourway_saved_ranges_20261001/README.md",
        "results/archcomp26_20261001/plots/nohash_saved_20261001/SUMMARY.md",
        "results/archcomp26_20261001/single_pendulum_two_state_p3_full20_001/SUMMARY.md",
        "results/archcomp26_20261001/p3_tora_remain_v1/full20_001/INDEPENDENT_INTERVAL_SCAN.json",
        "results/archcomp26_20261001/native_attitude_avoid_full30_001/SCAN.json",
        "results/archcomp26_20261001/native_attitude_avoid_full30_001/BOX_UNSAFE_AUDIT.json",
        "results/archcomp26_20261001/author_attitude_avoid_v1/huan_full30_001/payload/RESULT.json",
        "results/archcomp26_20261001/author_attitude_avoid_v1/xiangru_full30_001/payload/RESULT.json",
        "results/archcomp26_20261001/p3_attitude_avoid_v1/full30_001/RESULT.json",
        "results/archcomp26_20261001/DOCKING_FULLBOX_3METHODS_SUMMARY.md",
        "results/archcomp26_20261001/native_docking_full40_001/SUMMARY.md",
        "results/archcomp26_20261001/plots/nohash_saved_20261001/docking_fullbox_4method_q_upper.geometry.json",
        "results/archcomp26_20261001/airplane_continuous_order3_fullbox_20261002/SUMMARY.md",
        "results/archcomp26_20261001/AIRPLANE_P3_FULLBOX_SMOKES_20261002.md",
        "results/archcomp26_20261001/native_airplane_fullbox_smokes_20261002/SUMMARY.md",
        "results/huan_quad_stage_a_40_20261001/observer_pair_v1/SUMMARY.md",
        "output/flowstar_latest_20260930/repo/docs/ARCHCOMP26_ATTITUDE_CONTROL_CONTRACT_20261001.md",
    ):
        _insert(anchor, path, "Note")

    # Compact the provenance index so its last path does not spill alone.
    in_source_list = False
    for paragraph in doc.paragraphs:
        if paragraph.text == "完整报告与新结果":
            in_source_list = True
            continue
        if paragraph.text == "DP 绘图源与输出":
            break
        if in_source_list and paragraph.text.startswith(("results/", "output/")):
            paragraph.style = "Note"
            paragraph.paragraph_format.space_after = Pt(1)
            for run in paragraph.runs:
                run.font.size = Pt(7.5)

    # The cover already states the no-digest and unexecuted MATLAB scope;
    # repeating it as the final Appendix B line creates a page by itself in Word.
    redundant = _find(doc, "图、MATLAB 脚本和本报告在本轮均未进行内容摘要校验")
    redundant._element.getparent().remove(redundant._element)

    assert len(doc.inline_shapes) == 5
    assert "24 个方法单元" in _find(doc, "目标是在同一 benchmark 合同下").text
    assert all(row.cells[1].text == "全程" for row in coverage.rows[1:2])
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
